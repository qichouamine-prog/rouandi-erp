"""
excel_export.py - Génère des fichiers Excel EXPORTÉS avec une mise en forme
professionnelle (bandeau logo/titre coloré, en-têtes stylées, colonnes
alignées) au lieu d'un CSV brut sans habillage.

Utilise openpyxl (déjà listé dans requirements.txt via pandas[excel]).
"""

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FACTURE_DISPLAY_COLUMNS = {
    "mois": "Mois",
    "marche": "Marché / Chantier",
    "fournisseur": "Fournisseur",
    "nature": "Nature",
    "designation": "Désignation",
    "montant": "Montant (MAD)",
    "remarque": "Remarque",
    "contrat": "N° de Contrat",
    "num_facture": "N° de Facture",
    "facture_originale_recue": "Facture Originale ?",
    "recu_paiement_original_recu": "Reçu Original ?",
    "date_reglement": "Date de Règlement",
    "mode_reglement": "Mode de Règlement",
    "date_cheque": "Date de Chèque",
}

PRORATA_DISPLAY_COLUMNS = {
    "date": "Date",
    "chantier": "Chantier",
    "libelle": "Libellé Dépense",
    "fournisseur": "Fournisseur",
    "montant_total": "Montant Total (MAD)",
    "entreprise": "Entreprise / Sous-Traitant",
    "cle_repartition": "Clé de Répartition (%)",
    "montant_partage": "Montant Partagé (MAD)",
    "statut": "Statut Règlement",
}


def _styled_workbook(df, display_columns: dict, company_name: str, subtitle: str,
                      primary_color: str, secondary_color: str, sheet_title: str):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # limite Excel

    primary = primary_color.lstrip("#").upper() or "5B1B2D"
    secondary = secondary_color.lstrip("#").upper() or "C9A227"

    cols = [c for c in display_columns if c in df.columns]
    n_cols = max(len(cols), 1)

    # --- Bandeau titre société ---
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1, value=f"🏗  {company_name}")
    title_cell.font = Font(size=18, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=primary)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # --- Bandeau sous-titre (feuille + date de génération) ---
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub_cell = ws.cell(
        row=2, column=1,
        value=f"{sheet_title} — {subtitle} — généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}",
    )
    sub_cell.font = Font(size=10, italic=True, color="FFFFFF")
    sub_cell.fill = PatternFill("solid", fgColor=secondary)
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # ligne 3 vide (respiration visuelle)
    header_row = 4
    for j, key in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=j, value=display_columns[key])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=primary)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 24

    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        for j, key in enumerate(cols, start=1):
            val = row.get(key, "")
            if key in ("facture_originale_recue", "recu_paiement_original_recu"):
                val = "Oui" if bool(val) else "Non"
            cell = ws.cell(row=i, column=j, value=val)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F5F7")

    for j, key in enumerate(cols, start=1):
        if key in df.columns:
            col_series = df[key]
            if isinstance(col_series, pd.DataFrame):  # doublon de colonne résiduel
                col_series = col_series.iloc[:, 0]
            sample = col_series.fillna("").astype(str).tolist()[:300]
        else:
            sample = []
        max_len = max([len(str(display_columns[key]))] + [len(s) for s in sample] + [10])
        ws.column_dimensions[get_column_letter(j)].width = min(max_len + 3, 42)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    return wb


def export_factures_to_excel(df, company_name, company_subtitle, primary_color, secondary_color) -> io.BytesIO:
    wb = _styled_workbook(
        df, FACTURE_DISPLAY_COLUMNS, company_name, company_subtitle,
        primary_color, secondary_color, sheet_title="Suivi des Factures",
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_prorata_to_excel(df, company_name, company_subtitle, primary_color, secondary_color) -> io.BytesIO:
    wb = _styled_workbook(
        df, PRORATA_DISPLAY_COLUMNS, company_name, company_subtitle,
        primary_color, secondary_color, sheet_title="Compte Prorata",
    )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
